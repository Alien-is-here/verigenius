import streamlit as st
import zipfile
import tempfile
from pathlib import Path
import requests


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="VeriGenius",
    page_icon="🧪",
    layout="centered"
)


# ============================================================
# BACKEND CONFIG
# ============================================================

BACKEND_URL = st.secrets.get(
    "BACKEND_URL",
    "https://YOUR-BACKEND-URL/generate"
)


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_text(uploaded_file):

    extension = Path(
        uploaded_file.name
    ).suffix.lower()

    # --------------------------------------------------------
    # Markdown
    # --------------------------------------------------------

    if extension == ".md":

        return uploaded_file.read().decode(
            "utf-8"
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == ".docx":

        from docx import Document

        document = Document(
            uploaded_file
        )

        content = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                content.append(
                    paragraph.text.strip()
                )

        for table_index, table in enumerate(
            document.tables,
            1
        ):

            content.append(
                f"\n--- Table {table_index} ---"
            )

            for row in table.rows:

                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )

                if row_text.strip():

                    content.append(
                        row_text
                    )

        return "\n".join(content)

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        import PyPDF2

        reader = PyPDF2.PdfReader(
            uploaded_file
        )

        pages = []

        for page in reader.pages:

            text = page.extract_text()

            if text:

                pages.append(text)

        return "\n".join(pages)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


# ============================================================
# CREATE ZIP
# ============================================================

def create_zip(
    directory,
    zip_path
):

    directory = Path(
        directory
    )

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as archive:

        for file_path in directory.rglob("*"):

            if file_path.is_file():

                archive.write(
                    file_path,
                    file_path.relative_to(
                        directory
                    )
                )


# ============================================================
# HEADER
# ============================================================

st.title("VeriGenius")

st.subheader(
    "Generate a UVM Verification Blueprint"
)

st.write(
    "Upload your DUT specification and generate "
    "a structured UVM verification environment."
)


# ============================================================
# SPECIFICATION INPUT
# ============================================================

uploaded_spec = st.file_uploader(
    "Upload DUT Specification",
    type=[
        "pdf",
        "docx",
        "md"
    ],
    key="specification_file"
)

st.write("Or paste your specification:")

manual_spec = st.text_area(
    "Specification",
    height=220,
    key="manual_spec"
)


# ============================================================
# GENERATE
# ============================================================

if st.button(
    "Generate Blueprint",
    type="primary"
):

    specification = ""

    # --------------------------------------------------------
    # Read uploaded file
    # --------------------------------------------------------

    if uploaded_spec is not None:

        try:

            with st.spinner(
                "Reading specification..."
            ):

                specification = extract_text(
                    uploaded_spec
                )

        except Exception as e:

            st.error(
                f"Could not read specification: {e}"
            )

            st.stop()

    # --------------------------------------------------------
    # Read manually pasted specification
    # --------------------------------------------------------

    elif manual_spec.strip():

        specification = manual_spec

    else:

        st.warning(
            "Please upload a specification "
            "or paste one."
        )

        st.stop()

    # --------------------------------------------------------
    # Validate specification
    # --------------------------------------------------------

    if not specification.strip():

        st.error(
            "The specification contains no readable text."
        )

        st.stop()

    # --------------------------------------------------------
    # Save specification
    # --------------------------------------------------------

    st.session_state[
        "specification"
    ] = specification

    # ========================================================
    # CALL PRIVATE BACKEND
    # ========================================================

    with st.spinner(
        "Generating UVM environment..."
    ):

        try:

            response = requests.post(
                BACKEND_URL,
                json={
                    "specification": specification
                },
                timeout=600
            )

            response.raise_for_status()

            result = response.json()

            # ------------------------------------------------
            # Backend response
            # ------------------------------------------------

            if result.get("status") != "success":

                st.error(
                    result.get(
                        "message",
                        "Backend generation failed."
                    )
                )

                st.stop()

            st.session_state[
                "final_state"
            ] = result["result"]

        except requests.exceptions.RequestException as e:

            st.error(
                "Could not connect to VeriGenius backend."
            )

            st.code(
                str(e)
            )

            st.stop()

        except Exception as e:

            st.error(
                "Generation failed."
            )

            st.exception(e)

            st.stop()


# ============================================================
# RESULTS
# ============================================================

if "final_state" in st.session_state:

    final_state = st.session_state[
        "final_state"
    ]

    st.success(
        "UVM verification environment generated."
    )


    # ========================================================
    # GENERATED FILES
    # ========================================================

    st.subheader(
        "Generated UVM Files"
    )


    files = {

        "filled_boilerplate.sv":
            final_state.get(
                "filled_boilerplate",
                ""
            ),

        "structural_uvc.sv":
            final_state.get(
                "struct_code",
                ""
            ),

        "uvm_sequences.sv":
            final_state.get(
                "sequence_code",
                ""
            ),

        "uvm_tests.sv":
            final_state.get(
                "test_code",
                ""
            ),

        "symbol_registry.json":
            final_state.get(
                "symbol_registry",
                ""
            ),

        "gap_analysis.md":
            final_state.get(
                "gap_analysis_report",
                ""
            )
    }


    # --------------------------------------------------------
    # Remove empty files
    # --------------------------------------------------------

    files = {
        name: content
        for name, content in files.items()
        if content
    }


    if not files:

        st.warning(
            "The workflow returned no generated artifacts."
        )

    else:

        st.write(
            f"{len(files)} generated file(s)"
        )


        # ----------------------------------------------------
        # Individual files
        # ----------------------------------------------------

        for filename, content in files.items():

            with st.expander(
                filename
            ):

                if filename.endswith(".sv"):

                    st.code(
                        content,
                        language="systemverilog"
                    )

                elif filename.endswith(".json"):

                    st.code(
                        content,
                        language="json"
                    )

                else:

                    st.markdown(
                        content
                    )

                st.download_button(
                    label=f"Download {filename}",
                    data=content,
                    file_name=filename,
                    key=f"download_{filename}"
                )


        # ====================================================
        # ZIP DOWNLOAD
        # ====================================================

        temp_dir = Path(
            tempfile.mkdtemp()
        )

        output_dir = (
            temp_dir /
            "VeriGenius_Output"
        )

        output_dir.mkdir(
            parents=True,
            exist_ok=True
        )


        for filename, content in files.items():

            (
                output_dir /
                filename
            ).write_text(
                str(content),
                encoding="utf-8"
            )


        zip_path = (
            temp_dir /
            "VeriGenius_Output.zip"
        )


        create_zip(
            output_dir,
            zip_path
        )


        with open(
            zip_path,
            "rb"
        ) as zip_file:

            st.download_button(
                label="Download All Files as ZIP",
                data=zip_file,
                file_name="VeriGenius_Output.zip",
                mime="application/zip",
                type="primary",
                key="download_zip"
            )


    # ========================================================
    # USER FEEDBACK
    # ========================================================

    st.divider()

    st.subheader(
        "Submit Feedback"
    )

    st.write(
        "Review the generated UVM files above and "
        "describe any changes you want to make."
    )


    user_feedback = st.text_area(
        "Feedback",
        placeholder=(
            "Describe the changes you want.\n\n"
            "Example:\n"
            "Add a reset sequence.\n"
            "Add read/write corner-case tests.\n"
            "Modify the scoreboard comparison.\n"
            "Add coverage for FIFO full and empty conditions."
        ),
        height=180,
        key="user_feedback"
    )


    # ========================================================
    # SUBMIT FEEDBACK
    # ========================================================

    if st.button(
        "Submit Feedback",
        type="primary",
        key="submit_feedback_button"
    ):

        if not user_feedback.strip():

            st.warning(
                "Please enter feedback before submitting."
            )

        else:

            st.session_state[
                "submitted_feedback"
            ] = user_feedback.strip()

            st.success(
                "Feedback submitted successfully."
            )

            st.write(
                "Your feedback has been captured and is "
                "ready for the amendment stage."
            )

